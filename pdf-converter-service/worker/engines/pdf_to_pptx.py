import os
import glob
import subprocess
from pptx import Presentation
from pptx.util import Inches

def convert_pdf_to_pptx(input_pdf: str, output_pptx: str, dpi: int = 150) -> str:
    """
    Renders PDF pages at high fidelity and builds a presentation (.pptx).
    """
    temp_dir = os.path.dirname(output_pptx)
    prefix = os.path.join(temp_dir, "slide_page")

    # 1. Render pages to PNG using pdftoppm (C++ ultra-fast)
    subprocess.run(
        ["pdftoppm", "-png", "-r", str(dpi), input_pdf, prefix],
        check=True
    )

    page_images = sorted(glob.glob(f"{prefix}-*.png"))
    if not page_images:
        raise RuntimeError("Không thể kết xuất trang PDF sang slide PowerPoint.")

    # 2. Build PPTX
    prs = Presentation()
    # 16:9 widescreen dimensions (13.33 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    for img_path in page_images:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            img_path,
            left=0,
            top=0,
            width=prs.slide_width,
            height=prs.slide_height
        )

    prs.save(output_pptx)
    return output_pptx
