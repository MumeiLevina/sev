import os
import pdfplumber
import openpyxl

def convert_pdf_to_xlsx(input_pdf: str, output_xlsx: str) -> str:
    """
    Extracts all structured tables from PDF pages and saves to an Excel (.xlsx) workbook.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    table_found = False

    with pdfplumber.open(input_pdf) as pdf:
        for idx, page in enumerate(pdf.pages):
            tables = page.extract_tables()
            if not tables:
                # If no formal table found, extract words as lines
                text = page.extract_text()
                if text:
                    ws = wb.create_sheet(title=f"Trang_{idx + 1}")
                    for row_idx, line in enumerate(text.split('\n'), 1):
                        ws.cell(row=row_idx, column=1, value=line)
                    table_found = True
                continue

            for t_idx, table in enumerate(tables):
                sheet_title = f"Trang_{idx + 1}_Bang_{t_idx + 1}"[:31]
                ws = wb.create_sheet(title=sheet_title)
                for r_idx, row in enumerate(table, 1):
                    for c_idx, cell in enumerate(row, 1):
                        ws.cell(row=r_idx, column=c_idx, value=cell or "")
                table_found = True

    if not table_found:
        # Create at least an empty sheet with a note
        ws = wb.create_sheet(title="Thông báo")
        ws.cell(row=1, column=1, value="Không tìm thấy cấu trúc bảng trong tài liệu PDF này.")

    wb.save(output_xlsx)
    return output_xlsx
